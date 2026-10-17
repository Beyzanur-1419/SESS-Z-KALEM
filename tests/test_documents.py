from pathlib import Path

import fitz
import pytest
from docx import Document

import documents


def _bulgu(id_, kategori, kanit, baslangic, bitis, karar="kabul"):
    return {
        "id": id_,
        "kategori": kategori,
        "kanit": kanit,
        "konum": {"satir": 1, "baslangic": baslangic, "bitis": bitis},
        "karar": karar,
    }


# ---------- extract() ----------

def test_extract_txt(tmp_path):
    f = tmp_path / "belge.txt"
    f.write_text("Merhaba dünya", encoding="utf-8")
    assert documents.extract(f) == "Merhaba dünya"


def test_extract_pdf(tmp_path):
    f = tmp_path / "belge.pdf"
    pdf = fitz.open()
    sayfa = pdf.new_page()
    sayfa.insert_text((72, 72), "Hello synthetic contract")
    pdf.save(str(f))
    pdf.close()
    assert "Hello synthetic contract" in documents.extract(f)


def test_extract_docx_paragraflari_birlestirir(tmp_path):
    f = tmp_path / "belge.docx"
    belge = Document()
    belge.add_paragraph("Birinci satır")
    belge.add_paragraph("İkinci satır")
    belge.save(str(f))
    assert documents.extract(f) == "Birinci satır\nİkinci satır"


# ---------- _docx_haritasi() offsets ----------

def test_docx_haritasi_offsetleri_dogru(tmp_path):
    belge = Document()
    p = belge.add_paragraph()
    p.add_run("Toplam ")
    p.add_run("850.000 TL")
    p.add_run(" olarak ödenecektir.")

    metin, araliklar = documents._docx_haritasi(belge)
    assert metin == "Toplam 850.000 TL olarak ödenecektir."
    assert [(s, e, r.text) for s, e, r in araliklar] == [
        (0, 7, "Toplam "),
        (7, 17, "850.000 TL"),
        (17, 37, " olarak ödenecektir."),
    ]


# ---------- maskele_docx() ----------

def _paragraf_metni(dosya):
    return Document(str(dosya)).paragraphs[0].text


def test_maskele_docx_tek_run_tamamen_degisir(tmp_path):
    kaynak = tmp_path / "k.docx"
    cikti = tmp_path / "c.docx"
    belge = Document()
    p = belge.add_paragraph()
    p.add_run("Toplam ")
    p.add_run("850.000 TL")
    p.add_run(" olarak ödenecektir.")
    belge.save(str(kaynak))

    bulgular = [_bulgu("f1", "teklif_bedeli", "850.000 TL", 7, 17)]
    documents.maskele_docx(kaynak, cikti, bulgular)

    assert _paragraf_metni(cikti) == "Toplam [TEKLIF_BEDELI_1] olarak ödenecektir."


def test_maskele_docx_komsu_run_formati_korunur(tmp_path):
    kaynak = tmp_path / "k.docx"
    cikti = tmp_path / "c.docx"
    belge = Document()
    p = belge.add_paragraph()
    p.add_run("Toplam ")
    tutar = p.add_run("850.000 TL")
    tutar.bold = True
    p.add_run(" olarak ödenecektir.")
    belge.save(str(kaynak))

    bulgular = [_bulgu("f1", "teklif_bedeli", "850.000 TL", 7, 17)]
    documents.maskele_docx(kaynak, cikti, bulgular)

    sonuc = Document(str(cikti)).paragraphs[0]
    assert sonuc.text == "Toplam [TEKLIF_BEDELI_1] olarak ödenecektir."
    assert sonuc.runs[0].text == "Toplam "
    assert sonuc.runs[0].bold is not True
    assert sonuc.runs[-1].text == " olarak ödenecektir."
    assert sonuc.runs[-1].bold is not True


def test_maskele_docx_kismi_run_disindaki_metin_korunur(tmp_path):
    """Maskelenen değer bir run'ın yalnızca bir kısmını kaplıyorsa (aynı run
    içinde öncesi/sonrası var), run'ın geri kalanı ve formatı korunmalı."""
    kaynak = tmp_path / "k.docx"
    cikti = tmp_path / "c.docx"
    belge = Document()
    p = belge.add_paragraph()
    run = p.add_run("Bedel 850.000 TL burada.")
    run.italic = True
    belge.save(str(kaynak))

    baslangic = "Bedel ".__len__()
    bitis = baslangic + len("850.000 TL")
    bulgular = [_bulgu("f1", "teklif_bedeli", "850.000 TL", baslangic, bitis)]
    documents.maskele_docx(kaynak, cikti, bulgular)

    sonuc = Document(str(cikti)).paragraphs[0]
    assert sonuc.text == "Bedel [TEKLIF_BEDELI_1] burada."
    assert all(r.italic is True for r in sonuc.runs)


def test_maskele_docx_run_sinirini_asan_span_etiketi_bir_kez_yazar(tmp_path):
    """'Proje Kartal' iki ayrı run'a (normal + kalın) bölünmüş olsa bile
    etiket yalnızca bir kez yazılmalı, iki kez değil."""
    kaynak = tmp_path / "k.docx"
    cikti = tmp_path / "c.docx"
    belge = Document()
    p = belge.add_paragraph()
    p.add_run("Dahili olarak ")
    p.add_run("Proje ")
    kalin = p.add_run("Kartal")
    kalin.bold = True
    p.add_run(" adıyla anılmaktadır.")
    belge.save(str(kaynak))

    metin = documents.extract(kaynak)
    baslangic = metin.index("Proje Kartal")
    bitis = baslangic + len("Proje Kartal")
    bulgular = [_bulgu("f1", "kod_adi", "Proje Kartal", baslangic, bitis)]
    documents.maskele_docx(kaynak, cikti, bulgular)

    sonuc_metin = Document(str(cikti)).paragraphs[0].text
    assert sonuc_metin == "Dahili olarak [KOD_ADI_1] adıyla anılmaktadır."
    assert sonuc_metin.count("[KOD_ADI_1]") == 1


def test_maskele_docx_reddedilen_bulgu_degismez(tmp_path):
    kaynak = tmp_path / "k.docx"
    cikti = tmp_path / "c.docx"
    belge = Document()
    belge.add_paragraph("Proje Kartal devam ediyor.")
    belge.save(str(kaynak))

    bulgular = [_bulgu("f1", "kod_adi", "Proje Kartal", 0, 12, karar="red")]
    documents.maskele_docx(kaynak, cikti, bulgular)

    assert _paragraf_metni(cikti) == "Proje Kartal devam ediyor."


def test_maskele_docx_cok_paragrafli_belgede_dogru_paragrafi_bulur(tmp_path):
    kaynak = tmp_path / "k.docx"
    cikti = tmp_path / "c.docx"
    belge = Document()
    belge.add_paragraph("Birinci paragraf, gizli değer yok.")
    belge.add_paragraph("İkinci paragraf: TR330006100519786457841326 burada.")
    belge.save(str(kaynak))

    metin = documents.extract(kaynak)
    iban = "TR330006100519786457841326"
    baslangic = metin.index(iban)
    bitis = baslangic + len(iban)
    bulgular = [_bulgu("f1", "iban", iban, baslangic, bitis)]
    documents.maskele_docx(kaynak, cikti, bulgular)

    paragraflar = Document(str(cikti)).paragraphs
    assert paragraflar[0].text == "Birinci paragraf, gizli değer yok."
    assert paragraflar[1].text == "İkinci paragraf: [IBAN_1] burada."
