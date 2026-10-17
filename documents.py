"""Belge okuma (pdf/docx/txt) ve docx'e biçimlendirme koruyarak maskeli yazma.

bkz. docs/maskeleme-notu.md §3 — bu modülün var olma sebebi orada anlatılıyor:
tespit için metin, docx'in tüm paragraf+run'larının birleştirilmesiyle elde
edilir; bulgu offsetleri bu BİRLEŞİK metne göredir. Offsetler orijinal
run yapısına geri eşlenmeden yeni bir docx yazılırsa biçimlendirme
(kalın, italik, stil) tamamen kaybolur. Bu modül o eşlemeyi kurar ve
maskelemeyi doğrudan run'lar üzerinde, formatlarını koruyarak yapar.

Arayüz (Beyza, issue #6): `extract(dosya_yolu) -> str` — tek giriş noktası,
dosya tipine göre doğru okuyucuyu seçer.

Bilinen sınır: yalnızca `belge.paragraphs` taranır — tablo hücreleri içindeki
metin extraction'a dahil değildir (python-docx'te ayrı bir gezinme ister,
kapsam dışı bırakıldı, README'de not edilmeli).
"""

import copy
from pathlib import Path
from typing import Sequence, Tuple

import fitz  # PyMuPDF
from docx import Document
from docx.text.run import Run

import pseudonymize


def extract(dosya_yolu) -> str:
    """pdf, docx ya da txt dosyasından düz metin çıkarır."""
    yol = Path(dosya_yolu)
    uzanti = yol.suffix.lower()
    if uzanti == ".pdf":
        return _pdf_metni(yol)
    if uzanti == ".docx":
        metin, _ = _docx_haritasi(Document(str(yol)))
        return metin
    return yol.read_text(encoding="utf-8")


def _pdf_metni(yol: Path) -> str:
    with fitz.open(str(yol)) as pdf:
        return "\n".join(sayfa.get_text() for sayfa in pdf)


def _docx_haritasi(belge) -> Tuple[str, Sequence[Tuple[int, int, Run]]]:
    """(birleşik düz metin, [(baslangic, bitis, run_nesnesi), ...]) döner.

    Paragraflar "\\n" ile birleştirilir (extract() ile aynı kural); paragraf
    arası "\\n" karakterleri hiçbir run'a ait değildir. `run_nesnesi`,
    python-docx Run nesnesinin KENDİSİDİR (indeks değil) — maskeleme sırasında
    paragrafa yeni run'lar eklendiğinde indeks kayması oluşmaması için.
    """
    parcalar = []
    araliklar = []
    konum = 0
    paragraflar = belge.paragraphs
    for p_idx, paragraf in enumerate(paragraflar):
        for run in paragraf.runs:
            uzunluk = len(run.text)
            if uzunluk:
                araliklar.append((konum, konum + uzunluk, run))
                konum += uzunluk
            parcalar.append(run.text)
        if p_idx < len(paragraflar) - 1:
            parcalar.append("\n")
            konum += 1
    return "".join(parcalar), araliklar


def _run_sagina_ekle(run: Run, metin: str) -> Run:
    """`run` ile AYNI biçimlendirmeye sahip yeni bir run'ı hemen sağına ekler."""
    yeni_r = copy.deepcopy(run._r)
    run._r.addnext(yeni_r)
    yeni_run = Run(yeni_r, run._parent)
    yeni_run.text = metin
    return yeni_run


def maskele_docx(dosya_yolu, cikti_yolu, bulgular: Sequence[dict]) -> None:
    """Onaylı bulguları docx'e, biçimlendirmeyi KORUYARAK yazar.

    `bulgular`: docs/api-contract.md §2 şeması; `konum.baslangic`/`bitis`
    `_docx_haritasi()`'nin ürettiği BİRLEŞİK metne göre mutlak ofset olmalı.

    Bir bulgu birden çok run'a yayılıyorsa (ör. bir kelimenin yarısı kalın),
    etiket yalnızca İLK dokunulan run'da bir kez yazılır; sonraki run'larda
    o span'a denk gelen kısım yalnızca silinir (tekrar etiket yazılmaz).
    """
    belge = Document(str(dosya_yolu))
    _, araliklar = _docx_haritasi(belge)
    etiketler = pseudonymize.etiketle(bulgular)
    kabul = [b for b in bulgular if b.get("karar") == "kabul" and b.get("konum")]

    etiket_yayinlandi = set()
    for run_baslangic, run_bitis, run in araliklar:
        yerel = []
        for b in kabul:
            s, e = b["konum"]["baslangic"], b["konum"]["bitis"]
            if s < run_bitis and run_baslangic < e:
                local_s = max(0, s - run_baslangic)
                local_e = min(run_bitis - run_baslangic, e - run_baslangic)
                ilk_mi = b["id"] not in etiket_yayinlandi
                if ilk_mi:
                    etiket_yayinlandi.add(b["id"])
                yerel.append((local_s, local_e, ilk_mi, b["id"]))
        if not yerel:
            continue
        yerel.sort()
        orijinal_metin = run.text
        parcalar = []
        imlec = 0
        for local_s, local_e, ilk_mi, bid in yerel:
            parcalar.append(orijinal_metin[imlec:local_s])
            if ilk_mi:
                parcalar.append(f"[{etiketler[bid]}]")
            imlec = local_e
        parcalar.append(orijinal_metin[imlec:])

        run.text = parcalar[0]
        onceki = run
        for parca in parcalar[1:]:
            onceki = _run_sagina_ekle(onceki, parca)

    belge.save(str(cikti_yolu))
